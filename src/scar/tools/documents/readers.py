"""Document readers: PDF (pypdf), DOCX (python-docx), text formats, CSV/JSON/YAML, image metadata."""

from __future__ import annotations

import csv
import io
import json
from pathlib import Path
from typing import Any

import yaml

from scar.core.errors import ToolError

MAX_CHARS = 2_000_000


def read_pdf_bytes(data: bytes, max_pages: int = 500) -> str:
    from pypdf import PdfReader
    from pypdf.errors import PdfReadError

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception as exc:
                raise ToolError("the PDF is password-protected", "Encrypted") from exc
        parts: list[str] = []
        total = 0
        for i, page in enumerate(reader.pages[:max_pages]):
            text = page.extract_text() or ""
            parts.append(f"--- page {i + 1} ---\n{text.strip()}")
            total += len(text)
            if total > MAX_CHARS:
                break
        return "\n\n".join(parts)
    except PdfReadError as exc:
        raise ToolError(f"cannot read PDF: {exc}", "BadDocument") from exc


def read_pdf(path: Path) -> tuple[str, dict[str, Any]]:
    from pypdf import PdfReader

    data = path.read_bytes()
    text = read_pdf_bytes(data)
    meta: dict[str, Any] = {}
    try:
        reader = PdfReader(io.BytesIO(data))
        meta["pages"] = len(reader.pages)
        info = reader.metadata or {}
        meta.update({k.lstrip("/"): str(v) for k, v in dict(info).items() if isinstance(k, str)})
    except Exception as exc:  # noqa: BLE001 - metadata is optional; the text was already extracted
        meta["metadata_error"] = str(exc)[:120]
    if not text.replace("--- page", "").strip(" -\n0123456789"):
        meta["note"] = "no text layer (scanned PDF?) — use screen.ocr on page images"
    return text, meta


def read_docx(path: Path) -> tuple[str, dict[str, Any]]:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise ToolError(f"cannot read DOCX: {exc}", "BadDocument") from exc
    parts: list[str] = []
    for para in d.paragraphs:
        style = (para.style.name or "") if para.style is not None else ""
        if para.text.strip():
            prefix = "#" * int(style[-1]) + " " if style.startswith("Heading") and style[-1:].isdigit() else ""
            parts.append(prefix + para.text)
    for t_i, table in enumerate(d.tables):
        parts.append(f"\n[table {t_i + 1}]")
        for row in table.rows:
            parts.append(" | ".join(c.text.strip() for c in row.cells))
    cp = d.core_properties
    meta = {"title": cp.title, "author": cp.author, "paragraphs": len(d.paragraphs), "tables": len(d.tables)}
    return "\n".join(parts), meta


def read_text(path: Path, max_chars: int = MAX_CHARS) -> str:
    with path.open("r", encoding="utf-8-sig", errors="replace") as fh:
        return fh.read(max_chars)


def read_csv(path: Path, max_rows: int = 2000) -> tuple[str, dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", errors="replace", newline="") as fh:
        sample = fh.read(8192)
        fh.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample)
        except csv.Error:
            dialect = csv.excel
        rows = []
        for i, row in enumerate(csv.reader(fh, dialect)):
            if i >= max_rows:
                break
            rows.append(row)
    text = "\n".join(" | ".join(r) for r in rows)
    return text, {"rows_read": len(rows), "columns": len(rows[0]) if rows else 0, "header": rows[0] if rows else []}


def read_structured(path: Path) -> tuple[str, dict[str, Any]]:
    raw = read_text(path)
    try:
        data = json.loads(raw) if path.suffix.lower() == ".json" else yaml.safe_load(raw)
    except (json.JSONDecodeError, yaml.YAMLError) as exc:
        return raw, {"parse_error": str(exc)[:200]}
    kind = type(data).__name__
    size = len(data) if isinstance(data, list | dict) else None
    return raw, {"type": kind, "size": size}


def image_meta(path: Path) -> dict[str, Any]:
    from PIL import ExifTags, Image

    with Image.open(path) as img:
        meta: dict[str, Any] = {"format": img.format, "width": img.width, "height": img.height, "mode": img.mode}
        try:
            exif = img.getexif()
            for k, v in list(exif.items())[:40]:
                meta[str(ExifTags.TAGS.get(k, k))] = str(v)[:120]
        except Exception as exc:  # noqa: BLE001 - EXIF is optional
            meta["exif_error"] = str(exc)[:120]
    return meta


def read_any(path: Path) -> tuple[str, dict[str, Any]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return read_pdf(path)
    if ext == ".docx":
        return read_docx(path)
    if ext in (".csv", ".tsv"):
        return read_csv(path)
    if ext in (".json", ".yaml", ".yml"):
        return read_structured(path)
    if ext in (".png", ".jpg", ".jpeg", ".gif", ".bmp", ".webp", ".tif", ".tiff"):
        return "", image_meta(path)
    if ext in (".doc", ".xls", ".xlsx", ".ppt", ".pptx"):
        raise ToolError(f"{ext} files are not supported; save as PDF or DOCX", "Unsupported")
    return read_text(path), {}
