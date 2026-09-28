"""Document writers: Markdown/TXT, DOCX (python-docx), PDF (fpdf2) from lightweight Markdown."""

from __future__ import annotations

import re
from pathlib import Path


def _blocks(markdown: str) -> list[tuple[str, str]]:
    """Very small Markdown subset: headings, bullets, numbered items, code fences, paragraphs."""
    out: list[tuple[str, str]] = []
    in_code = False
    code: list[str] = []
    para: list[str] = []

    def flush() -> None:
        if para:
            out.append(("p", " ".join(s.strip() for s in para)))
            para.clear()

    for line in markdown.splitlines():
        if line.strip().startswith("```"):
            if in_code:
                out.append(("code", "\n".join(code)))
                code = []
            else:
                flush()
            in_code = not in_code
            continue
        if in_code:
            code.append(line)
            continue
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            flush()
            out.append((f"h{len(m.group(1))}", m.group(2).strip()))
        elif re.match(r"^\s*[-*+]\s+", line):
            flush()
            out.append(("li", re.sub(r"^\s*[-*+]\s+", "", line)))
        elif re.match(r"^\s*\d+[.)]\s+", line):
            flush()
            out.append(("ol", re.sub(r"^\s*\d+[.)]\s+", "", line)))
        elif not line.strip():
            flush()
        else:
            para.append(line)
    flush()
    if in_code and code:
        out.append(("code", "\n".join(code)))
    return out


def _strip_inline(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return re.sub(r"\[([^\]]+)\]\((https?://[^)]+)\)", r"\1 (\2)", text)


def write_docx(path: Path, title: str, markdown: str) -> None:
    import docx

    d = docx.Document()
    if title:
        d.add_heading(title, level=0)
    for kind, text in _blocks(markdown):
        if kind.startswith("h"):
            d.add_heading(_strip_inline(text), level=min(int(kind[1]), 4))
        elif kind == "li":
            d.add_paragraph(_strip_inline(text), style="List Bullet")
        elif kind == "ol":
            d.add_paragraph(_strip_inline(text), style="List Number")
        elif kind == "code":
            p = d.add_paragraph()
            run = p.add_run(text)
            run.font.name = "Consolas"
        else:
            d.add_paragraph(_strip_inline(text))
    d.save(str(path))


def write_pdf(path: Path, title: str, markdown: str) -> None:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    font = "Helvetica"
    uni = Path(r"C:\Windows\Fonts\arial.ttf")
    if uni.exists():
        pdf.add_font("Arial", "", str(uni))
        bold = Path(r"C:\Windows\Fonts\arialbd.ttf")
        pdf.add_font("Arial", "B", str(bold if bold.exists() else uni))
        font = "Arial"

    def safe(t: str) -> str:
        return t if font == "Arial" else t.encode("latin-1", errors="replace").decode("latin-1")

    width = pdf.w - pdf.l_margin - pdf.r_margin
    if title:
        pdf.set_font(font, "B", 18)
        pdf.multi_cell(width, 9, safe(title))
        pdf.ln(3)
    for kind, text in _blocks(markdown):
        t = safe(_strip_inline(text))
        if kind.startswith("h"):
            pdf.set_font(font, "B", {1: 16, 2: 14, 3: 12}.get(int(kind[1]), 11))
            pdf.ln(2)
            pdf.multi_cell(width, 7, t)
        elif kind in ("li", "ol"):
            pdf.set_font(font, "", 11)
            pdf.multi_cell(width, 6, ("• " if kind == "li" and font == "Arial" else "- ") + t)
        elif kind == "code":
            pdf.set_font("Courier", "", 9)
            pdf.multi_cell(width, 5, text.encode("latin-1", errors="replace").decode("latin-1"))
        else:
            pdf.set_font(font, "", 11)
            pdf.multi_cell(width, 6, t)
            pdf.ln(1)
    pdf.output(str(path))


def write_document(path: Path, title: str, markdown: str) -> None:
    ext = path.suffix.lower()
    if ext == ".docx":
        write_docx(path, title, markdown)
    elif ext == ".pdf":
        write_pdf(path, title, markdown)
    elif ext in (".md", ".markdown"):
        path.write_text((f"# {title}\n\n" if title else "") + markdown.strip() + "\n", encoding="utf-8")
    elif ext in (".txt", ".text", ""):
        plain = "\n".join(_strip_inline(t) if k != "code" else t for k, t in _blocks(markdown))
        path.write_text((f"{title}\n{'=' * len(title)}\n\n" if title else "") + plain + "\n", encoding="utf-8")
    else:
        raise ValueError(f"unsupported document type {ext}")
