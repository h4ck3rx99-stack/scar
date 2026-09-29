"""Chat with your files: a local full-text index of Documents, Downloads and Desktop.

Reads PDF, Word, PowerPoint, text/Markdown/CSV and code files, splits them into ~1200-character chunks and stores
them in %APPDATA%\\jarvis\\files.db (SQLite FTS5). Indexing runs in a background thread and only re-reads files
whose modification time changed.
"""
import logging
import os
import re
import sqlite3
import threading
import time
from pathlib import Path

from rapidfuzz import fuzz

from jarvis import config
from jarvis.skills.files import SKIP_DIRS, folder

log = logging.getLogger("docs")

TEXT_EXT = {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".h", ".html", ".css", ".sql", ".ipynb", ".tex", ".log"}
DOC_EXT = {".pdf", ".docx", ".pptx"}
MAX_BYTES = 40 * 1024 * 1024
MAX_FILES = 6000
CHUNK = 1200

_SCHEMA = """
CREATE TABLE IF NOT EXISTS files (path TEXT PRIMARY KEY, mtime REAL, name TEXT);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5(path UNINDEXED, name, text, tokenize='porter unicode61');
"""


def extract_text(p: Path) -> str:
    ext = p.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(p))
            return "\n".join((page.extract_text() or "") for page in reader.pages[:300])
        if ext == ".docx":
            import docx
            d = docx.Document(str(p))
            parts = [para.text for para in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text for c in row.cells))
            return "\n".join(parts)
        if ext == ".pptx":
            from pptx import Presentation
            out = []
            for i, slide in enumerate(Presentation(str(p)).slides, 1):
                texts = [sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text_frame.text]
                out.append(f"[Slide {i}] " + " / ".join(texts))
            return "\n".join(out)
        if ext == ".ipynb":
            import json
            nb = json.loads(p.read_text(encoding="utf-8", errors="ignore"))
            return "\n".join("".join(c.get("source", "")) for c in nb.get("cells", []))
        return p.read_text(encoding="utf-8", errors="ignore")
    except Exception as e:  # noqa: BLE001 - encrypted/corrupt files are skipped
        log.info("could not read %s: %s", p.name, type(e).__name__)
        return ""


def chunks(text: str):
    text = re.sub(r"[ \t]+", " ", text)
    paras = re.split(r"\n\s*\n|\n(?=[A-Z0-9•\-])", text)
    cur = ""
    for para in paras:
        if len(cur) + len(para) > CHUNK and cur:
            yield cur.strip()
            cur = ""
        cur += para + "\n"
        while len(cur) > CHUNK * 1.5:
            yield cur[:CHUNK].strip()
            cur = cur[CHUNK:]
    if cur.strip():
        yield cur.strip()


class FileIndex:
    def __init__(self, path=None):
        self.path = path or (config.APPDATA_DIR / "files.db")
        config.APPDATA_DIR.mkdir(parents=True, exist_ok=True)
        self._lock = threading.Lock()
        self.db = sqlite3.connect(str(self.path), check_same_thread=False)
        self.db.executescript(_SCHEMA)
        self.db.commit()
        self.indexed = 0
        self.busy = False

    def roots(self):
        return [p for p in (folder(n) for n in ("documents", "downloads", "desktop")) if p]

    def scan(self):
        """Index new/changed files, drop deleted ones. Safe to call repeatedly (e.g. every 30 minutes)."""
        self.busy = True
        t0 = time.time()
        try:
            with self._lock:
                known = dict(self.db.execute("SELECT path, mtime FROM files").fetchall())
            seen, n_new = set(), 0
            for root in self.roots():
                for dirpath, dirnames, filenames in os.walk(root):
                    dirnames[:] = [d for d in dirnames if d not in SKIP_DIRS and not d.startswith(".")]
                    for f in filenames:
                        p = Path(dirpath) / f
                        ext = p.suffix.lower()
                        if ext not in TEXT_EXT and ext not in DOC_EXT:
                            continue
                        try:
                            st = p.stat()
                        except OSError:
                            continue
                        if st.st_size > MAX_BYTES or (ext in TEXT_EXT and st.st_size > 2 * 1024 * 1024):
                            continue
                        key = str(p)
                        seen.add(key)
                        if known.get(key) == st.st_mtime:
                            continue
                        text = extract_text(p)
                        with self._lock:
                            self.db.execute("DELETE FROM chunks WHERE path=?", (key,))
                            self.db.executemany("INSERT INTO chunks(path, name, text) VALUES (?,?,?)",
                                                [(key, p.name, c) for c in chunks(text)][:400])
                            self.db.execute("INSERT OR REPLACE INTO files(path, mtime, name) VALUES (?,?,?)", (key, st.st_mtime, p.name))
                            self.db.commit()
                        n_new += 1
                        if len(seen) >= MAX_FILES:
                            break
            gone = set(known) - seen
            with self._lock:
                for key in gone:
                    self.db.execute("DELETE FROM chunks WHERE path=?", (key,))
                    self.db.execute("DELETE FROM files WHERE path=?", (key,))
                self.db.commit()
                self.indexed = self.db.execute("SELECT COUNT(*) FROM files").fetchone()[0]
            log.info("file index: %d files (%d updated, %d removed) in %.0fs", self.indexed, n_new, len(gone), time.time() - t0)
        finally:
            self.busy = False

    def search(self, query: str, limit: int = 5):
        words = [w for w in re.findall(r"[\w']+", query.lower()) if len(w) > 2]
        if not words:
            return []
        q = " OR ".join(f'"{w}"' for w in words[:12])
        with self._lock:
            rows = self.db.execute(
                "SELECT path, name, snippet(chunks, 2, '', '', ' ... ', 60) FROM chunks WHERE chunks MATCH ? ORDER BY bm25(chunks) LIMIT ?",
                (q, limit)).fetchall()
        return [{"file": name, "path": path, "text": snip} for path, name, snip in rows]

    def find_file(self, name: str):
        with self._lock:
            rows = self.db.execute("SELECT path, name FROM files").fetchall()
        best = max(rows, key=lambda r: fuzz.WRatio(name.lower(), r[1].lower()), default=None)
        if best and fuzz.WRatio(name.lower(), best[1].lower()) >= 70:
            return Path(best[0])
        return None

    def file_text(self, path: Path, limit=14000):
        with self._lock:
            rows = self.db.execute("SELECT text FROM chunks WHERE path=?", (str(path),)).fetchall()
        text = "\n".join(r[0] for r in rows) if rows else extract_text(path)
        return text[:limit]
